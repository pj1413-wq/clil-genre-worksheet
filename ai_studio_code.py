import io
import streamlit as st
from openai import OpenAI
from google import genai
from anthropic import Anthropic
from pypdf import PdfReader
from docx import Document

# 頁面標題與佈局
st.set_page_config(
    page_title="CLIL 雙語學科 Genre 鷹架生成器",
    page_icon="🎓",
    layout="wide"
)

st.title("🎓 CLIL 雙語學科：Genre 文本分析與國一講義學習單生成系統")
st.markdown("依據 **Dr. Angel M. Y. Lin (林敏宜教授)** 之 SFL Genre Pedagogy 設計，協助非英語專長學科教師快速降載備課。")

# 側邊欄設定
with st.sidebar:
    st.header("🔑 系統與學科設定")
    
    # 讓使用者任選 AI 廠商
    ai_provider = st.selectbox(
        "選擇 AI 廠商 (Provider)：",
        ["OpenAI", "Google Gemini", "Anthropic Claude"]
    )
    api_key = st.text_input(f"請輸入 {ai_provider} API Key：", type="password", help=f"需填入有效 {ai_provider} API Key 才能呼叫 AI 模型。")
    
    st.markdown("---")
    subject = st.selectbox(
        "學科領域 (Subject)：",
        ["自然科學 (Science)", "數學 (Mathematics)", "視覺藝術 (Visual Arts)", "家政 (Home Economics)", "健康與體育 (PE/Health)", "表演藝術 (Performing Arts)", "其他 (Other)"]
    )
    topic = st.text_input("單元主題 (Topic)：", value="Laboratory Safety Rules")
    
    st.markdown("---")
    st.info("""
    💡 **理論核心：**
    * **Register (語域)**: Field, Tenor, Mode
    * **Genre 階段**: Staged, goal-oriented
    * **國一適性化**: CEFR A1-A2、圖表化、雙語句型鷹架
    """)

# 主畫面：輸入英文文本
st.subheader("1. 提供學科英文文本")
tab_text, tab_file = st.tabs(["✍️ 直接貼上文字", "📂 上傳文件 (PDF / Word / TXT)"])

raw_text = ""

with tab_text:
    pasted_text = st.text_area(
        "請貼上欲進行教學的學科英文課文或材料：",
        height=180,
        placeholder="例：Keep the exits and passages clear. Wear proper goggles. Do not eat or drink in the lab..."
    )
    if pasted_text:
        raw_text = pasted_text

with tab_file:
    uploaded_file = st.file_uploader("支援 PDF, DOCX, TXT 檔案", type=["pdf", "docx", "txt"])
    if uploaded_file is not None:
        if uploaded_file.name.endswith(".txt"):
            raw_text = uploaded_file.read().decode("utf-8")
        elif uploaded_file.name.endswith(".pdf"):
            reader = PdfReader(uploaded_file)
            raw_text = "\n".join([page.extract_text() for page in reader.pages if page.extract_text()])
        elif uploaded_file.name.endswith(".docx"):
            doc = Document(uploaded_file)
            raw_text = "\n".join([p.text for p in doc.paragraphs])
        st.success(f"已成功載入檔案：{uploaded_file.name} (共 {len(raw_text)} 字)")

# 定義提示詞
SYSTEM_PROMPT = """
你是一位熟稔 Dr. Angel M. Y. Lin 的 CLIL 雙語教學與 SFL Genre 理論專家。
請協助台灣「非英語專長之國中學科教師」，將上傳的英文文本轉化為符合「台灣國一程度（CEFR A1-A2）」的雙語講義與學習單。

請嚴格按照以下四大區塊格式輸出（使用繁體中文為主要說明語言）：

### 【第一部分：教師備課專用 - SFL Register & Genre 分析】
1. **Field (場域/學科概念)**：本篇文本探討的學科內容與核心概念。
2. **Tenor (基調/互動關係)**：作者與讀者角色、語氣之正式度與情態（Modality）。
3. **Mode (語式/組織形式)**：語篇組織方式（如：點列式指令、說明文、圖文搭配）。
4. **Genre 類型與階段結構 (Stages & Phases)**：
   - 文本類型 (Genre Type): （例如：Procedural Text 指導型規章、Descriptive Report 描述報告等）
   - 各階段結構 (Stages & Phases): 標示出各階段之溝通目的與文本對應內容。

### 【第二部分：精選 3 大核心語言特徵 (針對國一學生挑選)】
- **特徵 1 (學科字彙 Field-specific Vocabulary)**：挑選 4-6 個核心詞彙，附自然發音/重音提點與繁體中文解釋。
- **特徵 2 (核心句型 Patterns)**：挑選 1-2 個最關鍵句型（如祈使句 Imperatives、情態助動詞 Can/Could、比較級等），列出 Sentence Frame。
- **特徵 3 (篇章功能語詞 Discourse Markers)**：標示順序詞、因果詞或條件詞（如：First, Next, If... then...）。

### 【第三部分：國一學生課堂講義 (Student Handout)】
- 排版簡明，專為國一學生設計。
- 包含「文本結構拆解表（Stages）」與「句型鷹架對照表（Sentence Patterns）」。
- 穿插適量繁體中文鷹架，降低學科內容因語言難度產生的負擔。

### 【第四部分：課堂隨堂學習單 (Classroom Worksheet)】
按照 Lin 教授的 4 階段課堂活動設計：
- **Activity 1: Structure Matching (階段辨識)**：功能與句子的配對練習。
- **Activity 2: Language Detective (語言偵探)**：圈選關鍵字或動詞。
- **Activity 3: Substitution Table (句型置換表)**：表格引導填空造句。
- **Activity 4: Guided Reconstruction (引導式仿寫產出)**：給予情境任務，讓學生運用前述特徵完成 2~3 句學科英文產出。
"""

# 生成按鈕
st.subheader("2. 執行分析與生成")
if st.button("🚀 開始進行 Genre 分析並生成教材", type="primary"):
    if not api_key:
        st.error(f"⚠️ 請先在左側欄位輸入 {ai_provider} API Key！")
    elif not raw_text.strip():
        st.warning("⚠️ 請先貼上文本或上傳文件！")
    else:
        user_message = f"學科領域：{subject}\n課程單元：{topic}\n文本內容：\n{raw_text}"
        
        with st.spinner("AI 正在以 SFL Genre 理論解構文本並打造國一專用鷹架教材，請稍候..."):
            try:
                # 根據使用者的選擇，分流呼叫不同廠商的 API
                if ai_provider == "OpenAI":
                    client = OpenAI(api_key=api_key)
                    response = client.chat.completions.create(
                        model="gpt-4o",
                        messages=[
                            {"role": "system", "content": SYSTEM_PROMPT},
                            {"role": "user", "content": user_message}
                        ],
                        temperature=0.3
                    )
                    output_result = response.choices[0].message.content

                elif ai_provider == "Google Gemini":
                    client = genai.Client(api_key=api_key)
                    response = client.models.generate_content(
                        model='gemini-2.5-flash',
                        contents=f"{SYSTEM_PROMPT}\n\n任務請求：\n{user_message}"
                    )
                    output_result = response.text

                elif ai_provider == "Anthropic Claude":
                    client = Anthropic(api_key=api_key)
                    message = client.messages.create(
                        model="claude-3-5-sonnet-latest",
                        max_tokens=4000,
                        system=SYSTEM_PROMPT,
                        messages=[
                            {"role": "user", "content": user_message}
                        ],
                        temperature=0.3
                    )
                    output_result = message.content[0].text
                
                st.session_state["result"] = output_result
                st.success("✅ 分析與教材生成完畢！")
                
            except Exception as e:
                st.error(f"執行出錯：{str(e)}，請確認您的 API Key 是否正確且有足夠額度。")

# 結果展示與下載區
if "result" in st.session_state:
    st.markdown("---")
    st.subheader("3. 教學成果預覽與下載")
    
    # 建立 Word 檔案供下載
    doc = Document()
    doc.add_heading(f"{subject} - {topic} CLIL Genre 教學教材", level=1)
    for line in st.session_state["result"].split("\n"):
        doc.add_paragraph(line)
    
    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    
    col1, col2 = st.columns([1, 4])
    with col1:
        st.download_button(
            label="📄 下載為 Word 講義 (.docx)",
            data=doc_io,
            file_name=f"{topic}_CLIL_Genre_Materials.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
    with col2:
        st.download_button(
            label="📝 下載為純文字檔 (.md)",
            data=st.session_state["result"],
            file_name=f"{topic}_CLIL_Genre_Materials.md",
            mime="text/markdown"
        )
        
    st.markdown(st.session_state["result"])
